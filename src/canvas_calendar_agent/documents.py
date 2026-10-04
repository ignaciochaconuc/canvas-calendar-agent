"""Selección y extracción local, determinística, de documentos académicos."""

from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import pymupdf
from openpyxl import load_workbook

from .models import (EventCandidate, ExtractedDocument, ExtractedSpreadsheet,
                     SpreadsheetRow)

FILE_KEYWORDS = ("programa", "calendario", "cronograma", "planificacion", "evaluacion",
                 "evaluaciones", "fechas", "proyecto", "syllabus", "schedule")
EVENT_WORDS = ("prueba", "examen", "interrogacion", "control", "entrega", "proyecto",
               "presentacion", "evaluacion", "certamen", "tarea", "fecha", "calendario",
               "semana", "quiz")
MONTHS = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
          "septiembre", "octubre", "noviembre", "diciembre")
DATE_RE = re.compile(r"\b(?:\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?|\d{4}-\d{2}-\d{2})\b")
PDF_MIMES = {"application/pdf", "application/x-pdf"}
XLSX_MIMES = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


def _plain(value: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", value.lower())
                   if unicodedata.category(c) != "Mn")


def is_relevant_file(file: dict[str, Any], *, max_bytes: int = 20 * 1024 * 1024) -> bool:
    name = str(file.get("display_name") or file.get("filename") or "")
    mime = str(file.get("content-type") or file.get("content_type") or "").lower()
    size = file.get("size", 0)
    supported = ((name.lower().endswith(".pdf") and mime in PDF_MIMES) or
                 (name.lower().endswith(".xlsx") and mime in XLSX_MIMES))
    return (supported
            and isinstance(size, (int, float)) and size <= max_bytes
            and any(re.search(rf"\b{re.escape(word)}\b", _plain(name)) for word in FILE_KEYWORDS))


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (date, time)):
        return value.isoformat(timespec="minutes") if isinstance(value, time) else value.isoformat()
    return str(value)


def extract_xlsx(path: Path, *, filename: str | None = None) -> ExtractedSpreadsheet:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    rows: list[SpreadsheetRow] = []
    try:
        for sheet in workbook.worksheets:
            for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                values = [_cell_text(value) for value in cells]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append(SpreadsheetRow(sheet.title, number, values))
        return ExtractedSpreadsheet(filename or path.name, next(iter(XLSX_MIMES)),
                                    workbook.sheetnames, rows)
    finally:
        workbook.close()


def spreadsheet_blocks(document: ExtractedSpreadsheet) -> list[tuple[str, list[int], str]]:
    blocks: list[tuple[str, list[int], str]] = []
    for sheet in document.sheet_names:
        rows = [row for row in document.rows if row.sheet_name == sheet]
        if not rows:
            continue
        header = rows[0]
        relevant = []
        for row in rows:
            normalized = _plain(" | ".join(row.values))
            if any(re.search(rf"\b{re.escape(word)}\b", normalized) for word in EVENT_WORDS):
                relevant.append(row)
            elif any(DATE_RE.search(value) for value in row.values):
                relevant.append(row)
        if relevant:
            selected = [header] + [row for row in relevant if row.row_number != header.row_number]
            text = f"Archivo: {document.filename}\nHoja: {sheet}\n\n" + "\n".join(
                " | ".join(row.values) for row in selected)
            blocks.append((sheet, [row.row_number for row in relevant], text))
    return blocks


def spreadsheet_candidates(document: ExtractedSpreadsheet, file: dict[str, Any],
                           course: dict[str, Any]) -> list[EventCandidate]:
    source_id = str(file.get("id", "unknown"))
    return [EventCandidate(int(course["id"]), str(course.get("name") or "Sin nombre"), "file",
                           f"{source_id}:sheet:{sheet}", f"{document.filename} — {sheet}", text,
                           file.get("url"), None,
                           {"file_type": "xlsx", "sheet": sheet, "rows": rows})
            for sheet, rows, text in spreadsheet_blocks(document)]


def extract_pdf(path: Path, *, filename: str | None = None) -> ExtractedDocument:
    pages: list[str] = []
    with pymupdf.open(path) as document:
        for page in document:
            text = page.get_text("text", sort=True).strip()
            tables = []
            try:
                for table in page.find_tables().tables:
                    rows = [" | ".join(str(cell or "").strip() for cell in row)
                            for row in table.extract()]
                    if rows:
                        tables.append("\n".join(rows))
            except Exception:
                pass
            pages.append("\n\n".join(part for part in (text, *tables) if part).strip())
    combined = "\n\n".join(f"--- Página {i} ---\n{text}" for i, text in enumerate(pages, 1) if text)
    return ExtractedDocument(filename or path.name, "application/pdf", len(pages), combined, pages)


def relevant_blocks(document: ExtractedDocument, *, max_chars: int = 6000) -> list[tuple[int, str]]:
    found: list[tuple[int, str]] = []
    for page_number, page in enumerate(document.pages, 1):
        chunks = [page[i:i + max_chars] for i in range(0, len(page), max_chars)]
        for chunk in chunks:
            normalized = _plain(chunk)
            has_event = any(word in normalized for word in EVENT_WORDS)
            has_date = bool(DATE_RE.search(normalized)) or any(month in normalized for month in MONTHS)
            if has_event and has_date:
                found.append((page_number, chunk.strip()))
    return found


def file_candidates(document: ExtractedDocument, file: dict[str, Any],
                    course: dict[str, Any]) -> list[EventCandidate]:
    source_id = str(file.get("id", "unknown"))
    published = file.get("updated_at") or file.get("created_at")
    try:
        published_at = datetime.fromisoformat(str(published).replace("Z", "+00:00")) if published else None
    except ValueError:
        published_at = None
    return [EventCandidate(int(course["id"]), str(course.get("name") or "Sin nombre"), "file",
                           f"{source_id}:page:{page}", f"{document.filename} — página {page}",
                           text, file.get("url"), published_at)
            for page, text in relevant_blocks(document)]
